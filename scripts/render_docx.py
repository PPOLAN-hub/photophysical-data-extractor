#!/usr/bin/env python3
"""Render a readable Word report from validated canonical PDE JSON."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import render_html


IDENTITY_FIELDS = ("compound", "host_matrix", "doping_ratio", "sample_state")
CONDITION_FIELDS = ("temperature", "atmosphere", "excitation", "delay", "gate_window")


def field_text(field, fallback="未报告"):
    return render_html.raw(field, fallback)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table):
    table.style = "Table Grid"
    table.autofit = True
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "EAF1F8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])


def set_field_label(cell, name):
    label = render_html.FIELD_LABELS.get(name, name)
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    position = 0
    for match in re.finditer(r"<sub>(.*?)</sub>", label):
        if match.start() > position:
            paragraph.add_run(re.sub(r"<[^>]+>", "", label[position : match.start()]))
        run = paragraph.add_run(match.group(1))
        run.font.subscript = True
        position = match.end()
    if position < len(label):
        paragraph.add_run(re.sub(r"<[^>]+>", "", label[position:]))


def add_key_value_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    style_table(table)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return table


def add_sample(document, sample):
    identity = sample.get("identity", {})
    conditions = sample.get("conditions", {})
    fields = sample.get("fields", {})
    compound = field_text(identity.get("compound"), "未命名化合物")
    row_id = str(sample.get("row_id") or "未标注")
    document.add_heading(f"{compound} · {row_id}", level=2)

    context_rows = []
    identity_names = list(IDENTITY_FIELDS) + sorted(set(identity) - set(IDENTITY_FIELDS))
    condition_names = list(CONDITION_FIELDS) + sorted(set(conditions) - set(CONDITION_FIELDS))
    for name in identity_names:
        value = field_text(identity.get(name), "")
        if value:
            context_rows.append((render_html.FIELD_LABELS.get(name, name), value))
    for name in condition_names:
        value = field_text(conditions.get(name), "")
        if value:
            context_rows.append((render_html.FIELD_LABELS.get(name, name), value))
    if context_rows:
        add_key_value_table(document, context_rows)

    metric_table = document.add_table(rows=1, cols=3)
    for cell, value in zip(metric_table.rows[0].cells, ("物理量", "报告值", "证据")):
        cell.text = value
    style_table(metric_table)
    field_names = list(render_html.METRIC_ORDER) + sorted(set(fields) - set(render_html.METRIC_ORDER))
    for name in field_names:
        field = fields.get(name)
        if not isinstance(field, dict) or field.get("status") != "reported":
            continue
        cells = metric_table.add_row().cells
        set_field_label(cells[0], name)
        cells[1].text = field_text(field)
        context = field.get("measurement_context")
        if isinstance(context, dict) and context:
            cells[1].add_paragraph("；".join(f"{key}={value}" for key, value in context.items()))
        cells[2].text = str(field.get("evidence_id") or "未关联")
    if len(metric_table.rows) == 1:
        cells = metric_table.add_row().cells
        cells[0].merge(cells[2]).text = "未报告目标物理量"

    if sample.get("notes"):
        document.add_paragraph(f"备注：{sample['notes']}")


def analysis_text(value):
    if isinstance(value, dict):
        return str(value.get("text") or "未提供")
    return str(value or "未提供")


def add_article(document, data, index, batch_mode):
    paper = data.get("paper", {})
    if batch_mode and index > 1:
        document.add_page_break()
    document.add_heading(str(paper.get("title") or f"PDE 文献 {index}"), level=1)
    metadata = [
        ("文献 ID", paper.get("paper_id") or "未报告"),
        ("DOI", paper.get("doi") or "未报告"),
        ("期刊", paper.get("journal") or paper.get("journal_year") or "未报告"),
        ("年份", paper.get("year") or "未报告"),
        ("作者", "、".join(map(str, paper.get("authors", []))) if isinstance(paper.get("authors"), list) else paper.get("authors") or "未报告"),
        ("引用格式", paper.get("official_citation") or paper.get("citation") or paper.get("publisher_style_citation") or "未报告"),
    ]
    add_key_value_table(document, metadata)

    analysis = data.get("article_analysis", {})
    if analysis:
        document.add_heading("文章要点", level=2)
        document.add_paragraph("一句话创新：" + analysis_text(analysis.get("one_sentence_innovation")))
        koi = analysis.get("koi", {})
        if isinstance(koi, dict):
            for key, value in koi.items():
                document.add_paragraph(f"{key}：{analysis_text(value)}")

    document.add_heading("样品与测量条件", level=1)
    for sample in data.get("samples", []):
        add_sample(document, sample)

    review = data.get("manual_review", [])
    document.add_heading(f"人工复核（{len(review)} 项）", level=1)
    if review:
        for item in review:
            document.add_paragraph(str(item), style="List Bullet")
    else:
        document.add_paragraph("未列出需人工复核项目。")

    ledger = data.get("evidence_ledger", [])
    document.add_heading(f"证据台账（{len(ledger)} 条）", level=1)
    evidence_table = document.add_table(rows=1, cols=5)
    for cell, value in zip(evidence_table.rows[0].cells, ("证据", "条件记录", "字段", "位置", "原文线索")):
        cell.text = value
    style_table(evidence_table)
    for evidence in ledger:
        cells = evidence_table.add_row().cells
        values = (
            evidence.get("evidence_id"),
            evidence.get("row_id"),
            evidence.get("field_name"),
            evidence.get("location"),
            evidence.get("quote"),
        )
        for cell, value in zip(cells, values):
            cell.text = str(value or "")


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(9)
    styles["Title"].font.name = "Microsoft YaHei"
    styles["Title"].font.color.rgb = RGBColor(35, 87, 165)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Microsoft YaHei"
        styles[name].font.color.rgb = RGBColor(35, 58, 87)


def main(source_text, output_text):
    source = Path(source_text)
    documents = render_html.load_documents(source)
    if not documents:
        raise ValueError("No paper data found in input JSON")
    document = Document()
    configure_document(document)
    document.core_properties.title = "PDE photophysical data report"
    document.core_properties.subject = "photophysical-data-extractor"
    document.core_properties.keywords = "PDE, evidence-traceable, photophysics"
    title = "PDE 批量光物理数据报告" if len(documents) > 1 else str(documents[0].get("paper", {}).get("title") or "PDE 光物理数据报告")
    heading = document.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("由已验证的 paper_data.json 确定性生成。")
    for index, data in enumerate(documents, start=1):
        add_article(document, data, index, len(documents) > 1)
    Path(output_text).parent.mkdir(parents=True, exist_ok=True)
    document.save(output_text)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: render_docx.py paper_data.json report.docx")
    main(sys.argv[1], sys.argv[2])
