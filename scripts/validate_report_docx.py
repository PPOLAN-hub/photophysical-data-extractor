#!/usr/bin/env python3
"""Verify that a PDE DOCX preserves canonical reported values and evidence IDs."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

import render_html


def document_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def main(json_path, docx_path):
    documents = render_html.load_documents(Path(json_path))
    report = Document(docx_path)
    text = document_text(report)
    problems = []
    if report.core_properties.subject != "photophysical-data-extractor":
        problems.append("DOCX renderer marker missing")
    for data in documents:
        paper = data.get("paper", {})
        for label, value in (("title", paper.get("title")), ("doi", paper.get("doi"))):
            if value and str(value) not in text:
                problems.append(f"missing paper {label}: {value}")
        for sample in data.get("samples", []):
            row_id = sample.get("row_id")
            if row_id and str(row_id) not in text:
                problems.append(f"missing sample row_id: {row_id}")
            for group in ("identity", "conditions", "fields"):
                for name, field in sample.get(group, {}).items():
                    if not isinstance(field, dict) or field.get("status") != "reported":
                        continue
                    value = field.get("raw_value")
                    evidence_id = field.get("evidence_id")
                    if value is not None and str(value) not in text:
                        problems.append(f"missing reported value {row_id}.{name}: {value}")
                    if evidence_id and str(evidence_id) not in text:
                        problems.append(f"missing evidence link {row_id}.{name}: {evidence_id}")
    if problems:
        for problem in problems:
            print(f"ERROR: {problem}")
        raise SystemExit(1)
    print("DOCX canonical-content validation passed")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: validate_report_docx.py paper_data.json report.docx")
    main(sys.argv[1], sys.argv[2])
